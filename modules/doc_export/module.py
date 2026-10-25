"""
Export a project document in another format (the Export button on the Documents tab).

One command (POST /modules/commands/doc-export, {"arg": "{\"doc_id\": 5, \"format\": \"docx\"}"}):
  formats: "docx" (Word), "txt" (plain text), "png" (first page of a PDF as an image, Mac only)
  from:    PDF -> docx, txt, png;  text-like files (md, txt, csv, json, html...) -> docx, txt
Returns {"name", "mime", "data"} with the file base64-encoded, or {"error": "<why, in words>"}.
(An exception would reach the page only as its type name, so problems are returned, not raised.)
Images are converted in the browser (canvas), so they never come here.
Needs pypdf, pdfminer.six and python-docx (requirements.txt); a missing one is reported, not a crash.

A module, so a change here needs POST /modules/doc_export/reload, never a restart.
"""

import base64
import io
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import db
import documents

TEXT = re.compile(r"\.(md|markdown|txt|csv|json|ya?ml|log|html?|xml|ts|tsx|js|jsx|py|sql|css)$", re.I)
MAX_BYTES = 25 * 1024 * 1024
MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "png": "image/png",
}


def _pdf_text(path: Path) -> list[str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ValueError("The PDF reader isn't installed on this computer (pip install pypdf)")
    pages = [(p.extract_text() or "").strip() for p in PdfReader(str(path)).pages]
    if not any(pages):
        raise ValueError("This PDF has no text in it (it's probably a scan), so there's nothing to turn into Word")
    return pages


def _docx(paragraphs: list[str]) -> bytes:
    try:
        from docx import Document
    except ImportError:
        raise ValueError("The Word writer isn't installed on this computer (pip install python-docx)")
    doc = Document()
    for i, page in enumerate(paragraphs):
        if i:
            doc.add_page_break()
        for block in re.split(r"\n\s*\n", page):
            if block.strip():
                doc.add_paragraph(block.strip())
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---- PDF -> Word keeping the layout: headings, bold/italic/colour, bullets, numbered lists, tables.
# Read with pdfminer.six (MIT) and rebuilt with python-docx. Pictures and drawn shapes are not carried over.
MARKER = re.compile(r"^(?:[•●▪◦·–—*-]|\d{1,3}[.)]|[a-zA-Z][.)])$")
FOOTER = re.compile(r"^(page\s+\d+(\s+of\s+\d+)?|\d+\s*/\s*\d+|\d+)$", re.I)


def _family(fontname: str) -> str:
    base = re.sub(r"^[A-Z]{6}\+", "", fontname or "")
    low = base.lower()
    if re.search(r"courier|mono|consolas|menlo", low):
        return "Courier New"
    if re.search(r"times|serif(?!.*sans)|georgia|garamond", low) and "sans" not in low:
        return "Times New Roman"
    if re.search(r"helvetica|arial", low):
        return "Arial"
    name = re.split(r"[-,]", base)[0]
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", name) or "Arial"


def _rgb(color):
    try:
        c = [float(x) for x in (color if isinstance(color, (tuple, list)) else [color])]
    except (TypeError, ValueError):
        return None
    if len(c) == 1:
        c = [c[0]] * 3
    elif len(c) == 4:
        k = c[3]
        c = [(1 - c[i]) * (1 - k) for i in range(3)]
    if len(c) != 3:
        return None
    rgb = tuple(max(0, min(255, int(round(v * 255)))) for v in c)
    return None if max(rgb) < 60 else rgb


def _line_runs(line):
    """[(text, bold, italic, size, rgb, family)] for one pdfminer text line."""
    from pdfminer.layout import LTChar
    runs: list[list] = []
    for ch in line:
        if isinstance(ch, LTChar):
            low = ch.fontname.lower()
            key = (any(w in low for w in ("bold", "black", "heavy")), any(w in low for w in ("italic", "oblique")),
                   round(ch.size * 2) / 2, _rgb(getattr(ch.graphicstate, "ncolor", None)), _family(ch.fontname))
            if runs and runs[-1][1:] == list(key):
                runs[-1][0] += ch.get_text()
            else:
                runs.append([ch.get_text()] + list(key))
        elif runs and getattr(ch, "get_text", None) and ch.get_text() == " ":
            runs[-1][0] += " "
    for r in runs:
        r[0] = r[0].replace("\n", "")
    return [tuple(r) for r in runs if r[0]]


def _pdf_layout_docx(path: Path) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Emu
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LAParams, LTRect, LTLine, LTTextContainer, LTTextLine

    doc = Document()
    first = True
    sizes: list[float] = []
    pages = []
    for page in extract_pages(str(path), laparams=LAParams(line_margin=0.2, boxes_flow=None)):
        lines, rules = [], []

        def walk(o):
            if isinstance(o, LTTextLine):
                lines.append(o)
            elif isinstance(o, (LTRect, LTLine)):
                if o.height < 2.5 and o.width > 6:
                    rules.append(o)
            elif hasattr(o, "__iter__"):
                for c in o:
                    walk(c)
        walk(page)
        pages.append((page, lines, rules))
        for ln in lines:
            for r in _line_runs(ln):
                sizes += [r[3]] * len(r[0])
    if not sizes:
        raise ValueError("This PDF has no text in it (it's probably a scan), so there's nothing to turn into Word")
    body = sorted(sizes)[len(sizes) // 2]

    for page, lines, rules in pages:
        W, H = page.width, page.height
        sec = doc.sections[0]
        if first:
            sec.page_width, sec.page_height = Emu(int(W * 12700)), Emu(int(H * 12700))
            xs = [ln.x0 for ln in lines] or [54]
            left = max(24.0, min(xs))
            right = max(24.0, W - max([ln.x1 for ln in lines] or [W - 54]))
            sec.left_margin, sec.right_margin = Emu(int(left * 12700)), Emu(int(right * 12700))
            sec.top_margin = sec.bottom_margin = Emu(int(min(60, max(36, H * 0.06)) * 12700))
        else:
            doc.add_page_break()
        first = False
        left = doc.sections[0].left_margin / 12700
        lines = [ln for ln in lines if ln.get_text().strip() and not (ln.y1 < H * 0.07 and FOOTER.match(ln.get_text().strip()))]
        # rows: lines sharing a baseline
        lines = [ln for ln in lines if re.search(r"\w|[\u2190-\u21ff]", re.sub(r"\(cid:\d+\)", "", ln.get_text())) or MARKER.match(ln.get_text().strip())]
        lines.sort(key=lambda ln: (-(ln.y0 + ln.y1) / 2, ln.x0))
        rows: list[list] = []
        for ln in lines:
            if rows and abs((rows[-1][0].y0 + rows[-1][0].y1) / 2 - (ln.y0 + ln.y1) / 2) < 5:
                rows[-1].append(ln)
            else:
                rows.append([ln])
        for r in rows:
            r.sort(key=lambda ln: ln.x0)

        def add_runs(par, runs):
            for text, bold, italic, size, rgb, fam in runs:
                run = par.add_run(text)
                run.bold, run.italic = bold or None, italic or None
                run.font.size, run.font.name = Pt(size), fam
                run.font.color.rgb = RGBColor(*(rgb or (0, 0, 0)))

        def underlined(ln):
            return any(r.x0 <= ln.x0 + 3 and r.x1 >= ln.x1 - 3 and r.x1 - r.x0 <= (ln.x1 - ln.x0) * 1.15 and -1 <= ln.y0 - r.y1 <= 5 for r in rules)

        # classify each row: list item, table row or text
        items = []
        for r in rows:
            segs = [_line_runs(ln) for ln in r]
            if len(r) >= 2 and MARKER.match(r[0].get_text().strip()) and r[1].x0 - r[0].x1 < 40:
                items.append(("li", r, segs))
            elif len(r) >= 2 and sum(r[i + 1].x0 - r[i].x1 > 40 for i in range(len(r) - 1)) >= 1 and r[-1].x0 - r[0].x0 > 150:
                items.append(("tr", r, segs))
            elif MARKER.match(r[0].get_text().strip().split(" ")[0]) and len(r[0].get_text().strip().split(" ")) > 1 and r[0].get_text().strip()[0] in "•●-–":
                items.append(("li1", r, segs))
            else:
                items.append(("tx", r, segs))

        i = 0
        prev = None
        while i < len(items):
            kind, r, segs = items[i]
            if kind == "tr":
                j = i
                grid = []
                while j < len(items) and items[j][0] == "tr":
                    grid.append(items[j])
                    j += 1
                cols = max(len(g[1]) for g in grid)
                if len(grid) >= 2 or cols >= 3:
                    tbl = doc.add_table(rows=0, cols=cols)
                    tbl.style = "Table Grid"
                    head = max(grid, key=lambda g: len(g[1]))
                    edges = [ln.x0 for ln in head[1]]
                    for g in grid:
                        cells = tbl.add_row().cells
                        for ln, seg in zip(g[1], g[2]):
                            k = max([e for e in range(len(edges)) if edges[e] <= ln.x0 + 6] or [0])
                            par = cells[k].paragraphs[0]
                            if par.runs:
                                par.add_run(" ")
                            add_runs(par, seg)
                    prev = grid[-1][1][0]
                    i = j
                    continue
            if kind in ("li", "li1"):
                if kind == "li":
                    marker = r[0].get_text().strip()
                    runs = [x for sg in segs[1:] for x in sg]
                    x_text = r[1].x0
                else:
                    first_run = segs[0][0]
                    marker, rest = first_run[0].strip().split(" ", 1)[0], first_run[0].strip().split(" ", 1)[1]
                    runs = [(rest,) + first_run[1:]] + segs[0][1:]
                    x_text = r[0].x0 + 14
                numbered = bool(re.match(r"^\w{1,3}[.)]$", marker)) and marker[0] not in "•●"
                par = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
                par.paragraph_format.left_indent = Pt(max(18, x_text - left))
                par.paragraph_format.first_line_indent = Pt(-14)
                par.paragraph_format.space_after = Pt(2)
                if numbered:
                    runs = [(marker + " ",) + runs[0][1:]] + runs
                    par.style = doc.styles["Normal"]
                    par.paragraph_format.left_indent = Pt(max(18, x_text - left))
                    par.paragraph_format.first_line_indent = Pt(-(x_text - r[0].x0))
                    runs[0] = (marker + "\t",) + runs[0][1:]
                add_runs(par, runs)
                last = r[-1]
                i += 1
                # continuation lines of this item: same text indent, close below
                while i < len(items) and items[i][0] == "tx" and abs(items[i][1][0].x0 - x_text) < 4 and 0 < last.y0 - items[i][1][0].y0 < body * 2.0:
                    add_runs(par, [(" ",) + runs[-1][1:]] + [x for x in items[i][2][0]])
                    last = items[i][1][0]
                    i += 1
                prev = last
                continue
            # plain text line(s) -> one paragraph
            runs = []
            for sg in segs:
                if runs and sg and not runs[-1][0].endswith(" "):
                    runs.append((" ",) + runs[-1][1:])
                runs += sg
            ln = r[0]
            size = max(x[3] for x in runs)
            heading = size >= body * 1.15 and all(x[1] for x in runs if x[0].strip())
            par = doc.add_paragraph()
            centered = abs((ln.x0 + r[-1].x1) / 2 - W / 2) < 12 and ln.x0 - left > 30
            if centered:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                par.paragraph_format.left_indent = Pt(max(0, ln.x0 - left)) if ln.x0 - left > 4 else None
            gap = (prev.y0 - ln.y1) if prev is not None else 0
            par.paragraph_format.space_before = Pt(max(0, min(24, gap - 2))) if prev is not None else Pt(0)
            par.paragraph_format.space_after = Pt(0)
            if heading:
                par.style = doc.styles["Heading 1" if size >= body * 1.5 else "Heading 2"]
                par.paragraph_format.space_before = Pt(max(0, min(24, gap - 2))) if prev is not None else Pt(0)
                par.paragraph_format.space_after = Pt(4)
            add_runs(par, runs)
            if underlined(ln):
                for run in par.runs:
                    run.underline = True
            last = r[-1]
            i += 1
            if not heading and not centered:
                # keep adding following lines of the same paragraph
                while i < len(items) and items[i][0] == "tx":
                    nr = items[i][1][0]
                    nsz = max(x[3] for x in items[i][2][0]) if items[i][2][0] else 0
                    if abs(nr.x0 - ln.x0) < 4 and 0 < last.y0 - nr.y0 < max(size, nsz) * 1.55 and abs(nsz - size) < 1.5:
                        add_runs(par, [(" ",) + runs[-1][1:]] + items[i][2][0])
                        last = items[i][1][0]
                        i += 1
                    else:
                        break
            prev = last
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _png(path: Path) -> bytes:
    if not shutil.which("sips"):
        raise ValueError("Turning a PDF into an image only works on a Mac for now")
    with tempfile.TemporaryDirectory() as tmp:
        dest = Path(tmp) / "page.png"
        r = subprocess.run(["sips", "-s", "format", "png", str(path), "--out", str(dest)],
                           capture_output=True, timeout=60)
        if r.returncode or not dest.exists():
            raise ValueError("Couldn't turn this PDF into an image")
        return dest.read_bytes()


def export(arg: str) -> dict:
    try:
        req = json.loads(arg or "{}")
        doc_id, fmt = int(req["doc_id"]), str(req["format"]).lower()
    except (ValueError, KeyError, TypeError):
        raise ValueError('Send {"doc_id": <id>, "format": "docx" | "txt" | "png"}')
    if fmt not in MIME:
        raise ValueError("Pick Word, text or image")
    doc = db.get_project_document(doc_id)
    path = documents.path_of(doc) if doc else None
    if path is None:
        raise ValueError("That document isn't there any more")
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("That file is too big to convert (over 25 MB)")
    is_pdf = doc["name"].lower().endswith(".pdf")
    if not is_pdf and not TEXT.search(doc["name"]):
        raise ValueError("This kind of file can't be converted yet")
    if is_pdf:
        if fmt == "png":
            data = _png(path)
        else:
            if fmt == "docx":
                try:
                    data = _pdf_layout_docx(path)
                except ImportError:
                    data = _docx(_pdf_text(path))
            else:
                data = "\n\n".join(_pdf_text(path)).encode("utf-8")
    else:
        if fmt == "png":
            raise ValueError("Text files can't be turned into an image")
        text = path.read_text(encoding="utf-8", errors="replace")
        data = _docx([text]) if fmt == "docx" else text.encode("utf-8")
    stem = Path(doc["name"]).stem
    return {"name": f"{stem}.{fmt}", "mime": MIME[fmt], "data": base64.b64encode(data).decode("ascii")}


class Module:
    def initialize(self, ctx):
        self.ctx = ctx
        ctx.register_command("doc-export", self.run, "Convert a project document to Word, text or an image")

    def run(self, arg: str) -> dict:
        try:
            return export(arg)
        except ValueError as e:
            return {"error": str(e)}
        except Exception as e:
            self.ctx.log.exception("doc export failed")
            return {"error": f"Couldn't convert this file ({type(e).__name__})"}
